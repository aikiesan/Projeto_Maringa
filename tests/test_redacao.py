"""
Testes do motor de redação (`tools.redacao`).

Tudo roda contra uma transcrição sintética construída aqui dentro, com nomes e
instituições fictícios. **Nenhum dado real do acervo entra em `tests/`** — é a
mesma regra que governa o resto do projeto.
"""

import os
import re
import sys
import unittest
import zipfile
from tempfile import TemporaryDirectory

sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from tools import redacao as R                                  # noqa: E402
from tools.transcript import turns                              # noqa: E402
from tools import supressoes as S                               # noqa: E402

PESSOA = "Joana Ribeiro Alves"
EMAIL = "joana.alves@exemplo.org"
FONE = "(44) 99999-0000"

LEXICO = [
    {"termo": "Instituto Fictício de Planejamento", "tipo": "nominal",
     "substituto": "[instituto de planejamento]", "escopo": "proprio", "nota": ""},
    {"termo": "IFP", "tipo": "acronimo",
     "substituto": "[instituto de planejamento]", "escopo": "proprio", "nota": ""},
    {"termo": "Sanefic", "tipo": "nominal",
     "substituto": "[companhia de saneamento]", "escopo": "sempre", "nota": ""},
    {"termo": "BNDES", "tipo": "acronimo",
     "substituto": "", "escopo": "nunca", "nota": "§3 terceiro público"},
]

PERMITIDOS = {R.dobrar(x) for x in
              ("Maringá", "Paraná", "Brasil", "BNDES", "Plano", "Diretor", "Agosto")}


def regras_base(proprias=("instituto fictício de planejamento", "ifp")):
    regras, _ = R.regras_pessoa(PESSOA, "[participante]")
    return (regras + R.regras_contato()
            + R.regras_instituicao(LEXICO, {R.dobrar(p) for p in proprias}))


class TestVariantes(unittest.TestCase):
    def test_gera_formas_esperadas(self):
        v = R.variantes_nome(PESSOA)
        self.assertIn("Joana Ribeiro Alves", v)
        self.assertIn("Joana Alves", v)
        self.assertIn("Joana", v)

    def test_particula_nao_vira_variante_solta(self):
        v = R.variantes_nome("Maria da Conceição Souza")
        self.assertNotIn("da", v)

    def test_token_comum_e_suprimido_e_relatado(self):
        _, supr = R.regras_pessoa("Carlos Silva", "[participante]")
        self.assertIn("Silva", supr)          # sobrenome comum não vira regra solta
        self.assertNotIn("Carlos", supr)


class TestSubstituicao(unittest.TestCase):
    def test_nome_completo_e_parciais(self):
        for entrada in ("Joana Ribeiro Alves", "Joana Alves", "Joana", "Alves"):
            saida, _ = R.aplicar(f"conversei com {entrada} ontem", regras_base())
            self.assertEqual(saida, "conversei com [participante] ontem", entrada)

    def test_tolerante_a_acento(self):
        regras, _ = R.regras_pessoa("José Antônio Pereira", "[participante]")
        saida, _ = R.aplicar("o Jose Antonio Pereira falou", regras)
        self.assertEqual(saida, "o [participante] falou")

    def test_caixa_alta_do_asr(self):
        """Regressão: «Cid» tem de casar «CID» — a transcrição automática grita siglas."""
        regras, _ = R.regras_pessoa("Cid Fictício", "[entrevistador]", apelidos=("Cid",))
        for grafia in ("Cid", "CID", "Cíd"):
            saida, _ = R.aplicar(f"passei pro {grafia} ontem", regras)
            self.assertEqual(saida, "passei pro [entrevistador] ontem", grafia)

    def test_token_solto_exige_maiuscula(self):
        """«alves» minúsculo no meio da frase não é nome — não pode ser redigido."""
        saida, _ = R.aplicar("ele alves pela janela", regras_base())
        self.assertEqual(saida, "ele alves pela janela")

    def test_apelido_curto_passa_pela_guarda(self):
        """Apelido explícito e curado vence a guarda de tamanho mínimo."""
        regras, _ = R.regras_pessoa("Ana Fictícia", "[participante]", apelidos=("Ana",))
        saida, _ = R.aplicar("a Ana comentou", regras)
        self.assertEqual(saida, "a [participante] comentou")

    def test_contato_antes_de_nome(self):
        """Regressão: e-mail tem de cair inteiro, não virar [participante].alves@…"""
        saida, _ = R.aplicar(f"manda pra {EMAIL}", regras_base())
        self.assertEqual(saida, "manda pra [e-mail]")
        self.assertNotIn("alves", saida)

    def test_identificadores(self):
        casos = {
            FONE: "[telefone]",
            "123.456.789-00": "[CPF]",
            "43.807.298/0001-59": "[CNPJ]",
            "https://exemplo.org/x": "[link]",
            "PIN: 227 064 004 7433#": "[PIN]",
        }
        for entrada, esperado in casos.items():
            saida, _ = R.aplicar(f"anota {entrada} ai", regras_base())
            self.assertIn(esperado, saida, entrada)

    def test_placeholder_nao_e_reescrito(self):
        """«[instituto de planejamento]» não pode ser tocado por uma regra de 'planejamento'."""
        extra = R.Regra(re.compile(r"planejamento", re.IGNORECASE), "[X]",
                        "instituicao", "manual", "planejamento")
        saida, _ = R.aplicar("trabalho no Instituto Fictício de Planejamento",
                             regras_base() + [extra])
        self.assertEqual(saida, "trabalho no [instituto de planejamento]")

    def test_escopo_proprio(self):
        texto = "o IFP cuida disso"
        propria, _ = R.aplicar(texto, regras_base())
        self.assertEqual(propria, "o [instituto de planejamento] cuida disso")
        alheia, _ = R.aplicar(texto, regras_base(proprias=()))
        self.assertEqual(alheia, texto)       # instituição de outra sessão fica nominal

    def test_escopo_sempre_e_nunca(self):
        saida, _ = R.aplicar("a Sanefic e o BNDES financiam", regras_base())
        self.assertIn("[companhia de saneamento]", saida)
        self.assertIn("BNDES", saida)         # §3: terceiro público permanece nominal


class TestResiduos(unittest.TestCase):
    def test_sinaliza_desconhecido_e_ignora_allowlist(self):
        achados = R.residuos("fomos a Curitiba discutir com Maringá", PERMITIDOS)
        tokens = {r.token for r in achados}
        self.assertIn("Curitiba", tokens)
        self.assertNotIn("Maringá", tokens)

    def test_nao_sinaliza_dentro_de_placeholder(self):
        achados = R.residuos("o [instituto de planejamento] respondeu", PERMITIDOS)
        self.assertEqual([r.token for r in achados], [])

    def test_sigla_e_risco_alto(self):
        achados = R.residuos("mandaram pro CREA ontem", PERMITIDOS)
        self.assertEqual([r.risco for r in achados if r.token == "CREA"], ["alto"])

    def test_cargo_eleva_risco(self):
        achados = R.residuos("falei com o secretário Fulaninho sobre isso", PERMITIDOS)
        alto = [r.token for r in achados if r.risco == "alto"]
        self.assertIn("Fulaninho", alto)


class TestVarreduraReversa(unittest.TestCase):
    """A verificação final tem de reprovar texto mal redigido — golden negative."""

    def test_detecta_vazamento_plantado(self):
        vazou = f"o {PESSOA} disse que sim"
        self.assertIn(R.dobrar("Alves"), R.dobrar(vazou))
        limpo, _ = R.aplicar(vazou, regras_base())
        self.assertNotIn(R.dobrar("alves"), R.dobrar(limpo))
        self.assertNotIn(R.dobrar("joana"), R.dobrar(limpo))


def _docx_sintetico(caminho):
    """Escreve um .docx mínimo no formato que `tools.transcript` lê."""
    paras = ["ago. 7, 2026", "Entrevista fictícia - Transcrição", "00:04:35",
             f"Cid Fictício: bom dia, {PESSOA}",
             f"{PESSOA}: bom dia, meu e-mail é {EMAIL}",
             "00:12:51",
             f"{PESSOA}: aqui no Instituto Fictício de Planejamento a gente cuida disso"]
    corpo = "".join(f"<w:p><w:r><w:t>{p}</w:t></w:r></w:p>" for p in paras)
    doc = ('<?xml version="1.0"?><w:document '
           'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           f"<w:body>{corpo}</w:body></w:document>")
    with zipfile.ZipFile(caminho, "w") as z:
        z.writestr("word/document.xml", doc)


class TestIntegracao(unittest.TestCase):
    def test_ponta_a_ponta(self):
        with TemporaryDirectory() as tmp:
            caminho = os.path.join(tmp, "ENT-999_Transcrição.docx")
            _docx_sintetico(caminho)
            tt = turns(caminho)
            rotulos = {"Cid Fictício": "[entrevistador]", PESSOA: "[participante]"}
            saida, subs, res = R.varrer_turnos(tt, regras_base(), PERMITIDOS, rotulos)

            self.assertEqual([t["ts"] for t in saida],
                             ["00:04:35", "00:04:35", "00:12:51"])
            texto = " ".join(t["texto"] for t in saida)
            self.assertNotIn("Joana", texto)
            self.assertNotIn("exemplo.org", texto)
            self.assertIn("[instituto de planejamento]", texto)
            self.assertTrue(subs)
            self.assertEqual({t["rotulo"] for t in saida},
                             {"[entrevistador]", "[participante]"})

    def test_timestamps_preservados(self):
        with TemporaryDirectory() as tmp:
            caminho = os.path.join(tmp, "ENT-999_Transcrição.docx")
            _docx_sintetico(caminho)
            originais = {t.ts for t in turns(caminho)}
            rotulos = {"Cid Fictício": "[entrevistador]", PESSOA: "[participante]"}
            saida, _, _ = R.varrer_turnos(turns(caminho), regras_base(),
                                          PERMITIDOS, rotulos)
            self.assertEqual({t["ts"] for t in saida}, originais)




# --------------------------------------------------------------- supressão
class TestSupressao(unittest.TestCase):
    """Supressão de trecho autoidentificador (`tools.supressoes`).

    Dados sintéticos, como o resto do arquivo. O que se verifica é o que a
    proteção promete: o trecho sai, a marca entra, o resto do turno sobrevive,
    e uma âncora que não casa **reprova** em vez de passar em silêncio.
    """

    TEXTO = ("[participante 1]: Eu presidi a comissão de meio ambiente da Câmara "
             "em 2019, e por isso conheço o processo.")
    TRECHO = "Eu presidi a comissão de meio ambiente da Câmara em 2019, e "

    def nova(self, **kw):
        campos = dict(fonte="transcricao", alvo="ENT-999", paragrafo=7,
                      texto=self.TEXTO, trecho=self.TRECHO,
                      substituto=S.MARCA,
                      motivo="presidencia nominal de comissao identifica a pessoa")
        campos.update(kw)
        return S.nova(**campos)

    def test_trecho_sai_e_marca_entra(self):
        """`substituto` com a marca: o corte fica visível, que é a convenção."""
        sup = self.nova()
        saida = S.aplicar_em(self.TEXTO, sup)
        self.assertNotIn("presidi a comissão", saida)
        self.assertNotIn("2019", saida)
        self.assertIn(S.MARCA, saida)

    def test_substituto_vazio_remove_sem_marca(self):
        """`substituto` vazio: o trecho sai e nada entra.

        É o tratamento certo quando o que identifica é o marcador de primeira
        pessoa e não o conteúdo: «o contrato aqui no IAM» perde só o «aqui no
        IAM» e a frase segue inteira. A coluna é literal justamente para que
        este caso seja expressável.
        """
        sup = self.nova(substituto="")
        saida = S.aplicar_em(self.TEXTO, sup)
        self.assertNotIn("presidi a comissão", saida)
        self.assertNotIn(S.MARCA, saida)
        self.assertIn("conheço o processo", saida)

    def test_substituto_com_texto_troca_o_trecho(self):
        """`substituto` com texto: substituição, como `redacao.py` faz com nome."""
        sup = self.nova(substituto="participei da comissão e ")
        saida = S.aplicar_em(self.TEXTO, sup)
        self.assertIn("participei da comissão e ", saida)
        self.assertNotIn("presidi", saida)
        self.assertNotIn("Câmara", saida)

    def test_trecho_no_fim_do_paragrafo(self):
        """Sufixo vazio é borda legítima, não declaração malformada.

        Um trecho que termina o parágrafo não tem o que vir depois dele.
        Exigir prefixo e sufixo não vazios fazia esses casos nunca casarem, e o
        sintoma era uma supressão declarada que a geração recusava aplicar.
        """
        texto = "[participante 1]: avançando um pouco, eu sou do Instituto X"
        trecho = "eu sou do Instituto X"
        sup = S.nova("transcricao", "ENT-999", 3, texto, trecho,
                     "vinculo institucional no fim do paragrafo",
                     substituto=S.MARCA)
        self.assertEqual(sup.sufixo, "")
        saida = S.aplicar_em(texto, sup)
        self.assertNotIn("Instituto X", saida)
        self.assertTrue(saida.endswith(S.MARCA))

    def test_resto_do_turno_sobrevive(self):
        saida = S.aplicar_em(self.TEXTO, self.nova())
        self.assertTrue(saida.startswith("[participante 1]: "))
        self.assertIn("conheço o processo", saida)

    def test_declaracao_nao_guarda_o_texto_suprimido(self):
        """O codebook é versionado: o material protegido não pode morar nele."""
        sup = self.nova()
        serializado = " ".join(str(getattr(sup, c)) for c in S.CAMPOS)
        self.assertNotIn("presidi", serializado)
        self.assertNotIn("Câmara", serializado)
        self.assertEqual(sup.n_chars, len(self.TRECHO))

    def test_ancora_que_nao_casa_reprova(self):
        """O cenário que o módulo existe para impedir: regerar e publicar sem o
        corte, sem ninguém perceber."""
        sup = self.nova()
        outro = self.TEXTO.replace("comissão", "subcomissão")
        with self.assertRaises(S.SupressaoNaoAplicada):
            S.aplicar_em(outro, sup)

    def test_trecho_alterado_reprova_mesmo_com_contexto_intacto(self):
        sup = self.nova()
        adulterado = self.TEXTO.replace("2019", "2020")
        with self.assertRaises(S.SupressaoNaoAplicada):
            S.aplicar_em(adulterado, sup)

    def test_aplica_nos_turnos_pelo_indice(self):
        sup = self.nova()
        corpo = [{"idx": 6, "ts": "00:00:10", "rotulo": "[participante 1]",
                  "texto": "Turno anterior, intacto."},
                 {"idx": 7, "ts": "00:00:20", "rotulo": "[participante 1]",
                  "texto": self.TEXTO}]
        saida, n = S.aplicar_em_turnos(corpo, "ENT-999", [sup])
        self.assertEqual(n, 1)
        self.assertEqual(saida[0]["texto"], "Turno anterior, intacto.")
        self.assertIn(S.MARCA, saida[1]["texto"])
        self.assertEqual(saida[1]["ts"], "00:00:20")

    def test_turno_ausente_reprova(self):
        sup = self.nova(paragrafo=999)
        corpo = [{"idx": 7, "ts": "00:00:20", "rotulo": "[participante 1]",
                  "texto": self.TEXTO}]
        with self.assertRaises(S.SupressaoNaoAplicada):
            S.aplicar_em_turnos(corpo, "ENT-999", [sup])

    def test_tolerante_aplica_o_que_falta(self):
        """A camada publicada recebe material que JA passou pelo pipeline."""
        sup = self.nova()
        pars = ["outro turno", "", self.TEXTO]
        sup = self.nova(paragrafo=2)
        saida, novas, ja = S.aplicar_tolerante(pars, "transcricao", "ENT-999", [sup])
        self.assertEqual((novas, ja), (1, 0))
        self.assertIn(S.MARCA, saida[2])

    def test_tolerante_aceita_o_que_ja_estava_aplicado(self):
        """O Anexo 05 já traz os cortes antigos escritos como (…).

        Exigir que essas âncoras casem de novo faria a geração falhar sempre.
        """
        sup = self.nova(paragrafo=0)
        pars = ["[participante 1]: " + S.MARCA + "por isso conheço o processo."]
        saida, novas, ja = S.aplicar_tolerante(pars, "transcricao", "ENT-999", [sup])
        self.assertEqual((novas, ja), (0, 1))
        self.assertEqual(saida, pars)

    def test_tolerante_reprova_quando_nem_casa_nem_esta_aplicado(self):
        """A tolerância é estreita: fora dos dois casos, levanta.

        Âncora que não casa e não está aplicada é proteção perdida em silêncio,
        que é o desfecho que este módulo existe para impedir.
        """
        sup = self.nova(paragrafo=0)
        pars = ["[participante 1]: um turno qualquer, sem o trecho declarado."]
        with self.assertRaises(S.SupressaoNaoAplicada):
            S.aplicar_tolerante(pars, "transcricao", "ENT-999", [sup])

    def test_sessao_sem_supressao_passa_intacta(self):
        corpo = [{"idx": 7, "ts": "00:00:20", "rotulo": "[participante 1]",
                  "texto": self.TEXTO}]
        saida, n = S.aplicar_em_turnos(corpo, "ENT-000", [self.nova()])
        self.assertEqual(n, 0)
        self.assertEqual(saida[0]["texto"], self.TEXTO)


class _Run:
    def __init__(self, texto):
        self.text = texto


class _Par:
    """Dublê de parágrafo do python-docx.

    `tools.supressoes` não importa `docx`: exige de um parágrafo apenas `.text`,
    `.runs` e `.add_run`. O dublê existe para que a regra de supressão seja
    testável sem depender da biblioteca nem de um arquivo no disco.
    """

    def __init__(self, *textos):
        self.runs = [_Run(t) for t in textos]

    @property
    def text(self):
        return "".join(r.text for r in self.runs)

    def add_run(self, texto):
        self.runs.append(_Run(texto))
        return self.runs[-1]


class TestSupressaoEmParagrafo(unittest.TestCase):
    """`aplicar_em_paragrafo`: o corte dentro do .docx, sem perder formatação.

    O parágrafo real é multi-run, e reescrevê-lo inteiro apagaria o negrito do
    rótulo que abre cada afirmação da Seção 3. Aqui cada run é uma string, e o
    que se verifica é que só os runs atravessados pelo trecho mudam.
    """

    ABRE = "[participante 1]: "
    ALVO = "Eu presidi a comissão de meio ambiente da Câmara em 2019, e "
    FECHA = "por isso conheço o processo."

    def nova(self, **kw):
        campos = dict(fonte="transcricao", alvo="ENT-999", paragrafo=0,
                      texto=self.ABRE + self.ALVO + self.FECHA,
                      trecho=self.ALVO, substituto=S.MARCA,
                      motivo="presidencia nominal de comissao identifica a pessoa")
        campos.update(kw)
        return S.nova(**campos)

    def par(self):
        return _Par(self.ABRE, self.ALVO, self.FECHA)

    def test_runs_fora_do_trecho_ficam_intactos(self):
        par = self.par()
        self.assertTrue(S.aplicar_em_paragrafo(par, self.nova()))
        self.assertEqual(len(par.runs), 3)
        self.assertEqual(par.runs[0].text, self.ABRE)
        self.assertEqual(par.runs[2].text, self.FECHA)
        self.assertNotIn("presidi", par.text)

    def test_substituto_literal_entra_no_lugar_da_marca(self):
        """O defeito que este teste fixa: o aplicador inseria `MARCA` fixo.

        Cinco das supressões de transcrição declaram substituto que não é a
        marca. Com o defeito, o .docx de entrega saía divergente da página
        publicada justamente nesses cinco pontos.
        """
        sup = self.nova(substituto="participei da comissão e ")
        par = self.par()
        self.assertTrue(S.aplicar_em_paragrafo(par, sup))
        self.assertIn("participei da comissão e ", par.text)
        self.assertNotIn(S.MARCA, par.text)
        self.assertNotIn("Câmara", par.text)

    def test_equivale_ao_aplicador_de_texto(self):
        """Arquivo e página têm de produzir o mesmo texto, por construção."""
        for substituto in (S.MARCA, "", "participei da comissão e "):
            with self.subTest(substituto=substituto):
                sup = self.nova(substituto=substituto)
                par = self.par()
                S.aplicar_em_paragrafo(par, sup)
                self.assertEqual(par.text, S.aplicar_em(self.ABRE + self.ALVO + self.FECHA, sup))

    def test_ancora_que_nao_casa_devolve_falso(self):
        par = _Par("[participante 1]: um turno sem o trecho declarado.")
        self.assertFalse(S.aplicar_em_paragrafo(par, self.nova()))

    def test_aplicar_em_doc_usa_o_indice_cru(self):
        """Os índices declarados são os de `doc.paragraphs`, vazios inclusive."""
        sup = self.nova(paragrafo=2)
        pars = [_Par("cabeçalho"), _Par(""), self.par(), _Par("depois")]
        novas, ja = S.aplicar_em_doc(pars, "transcricao", "ENT-999", [sup])
        self.assertEqual((novas, ja), (1, 0))
        self.assertNotIn("presidi", pars[2].text)
        self.assertEqual(pars[3].text, "depois")


class TestJaAplicada(unittest.TestCase):
    """`_ja_aplicada`: reconhecer o tratamento sem afrouxar a proteção."""

    TEXTO = ("[participante 1]: o contrato aqui no Instituto Fictício de "
             "Planejamento venceu em 2019.")
    TRECHO = "aqui no Instituto Fictício de Planejamento "

    def nova(self, substituto=""):
        return S.nova(fonte="transcricao", alvo="ENT-999", paragrafo=0,
                      texto=self.TEXTO, trecho=self.TRECHO,
                      substituto=substituto,
                      motivo="local de trabalho do participante")

    def test_trecho_ainda_presente_nao_conta_como_aplicada(self):
        """O teste negativo que importa.

        Se um parágrafo intocado passasse por «já tratado», a proteção sumiria
        em silêncio, que é o desfecho que o módulo inteiro existe para impedir.
        """
        self.assertFalse(S._ja_aplicada(self.TEXTO, self.nova()))
        self.assertFalse(S._ja_aplicada(self.TEXTO, self.nova(S.MARCA)))

    def test_substituto_vazio_reconhece_contextos_adjacentes(self):
        """Remoção sem marca não deixa o que procurar.

        A prova positiva é que prefixo e sufixo ficaram colados. Sem isso,
        reprocessar um arquivo já corrigido fazia a geração falhar duro sobre
        material que estava certo.
        """
        sup = self.nova()
        tratado = S.aplicar_em(self.TEXTO, sup)
        self.assertNotIn(S.MARCA, tratado)
        self.assertTrue(S._ja_aplicada(tratado, sup))

    def test_tolerante_aceita_remocao_sem_marca(self):
        sup = self.nova()
        tratado = S.aplicar_em(self.TEXTO, sup)
        _, novas, ja = S.aplicar_tolerante([tratado], "transcricao", "ENT-999",
                                           [sup])
        self.assertEqual((novas, ja), (0, 1))

    def test_substituto_com_texto_e_prova_direta(self):
        sup = self.nova("no órgão ")
        tratado = S.aplicar_em(self.TEXTO, sup)
        self.assertTrue(S._ja_aplicada(tratado, sup))


class TestPremissaProduto4(unittest.TestCase):
    def test_supressoes_do_produto4_usam_a_marca(self):
        """Premissa da correção do `substituto`, fixada para caducar com aviso.

        As duas declarações de `fonte=produto4` têm substituto igual à marca, e
        por isso trocar `MARCA` fixo por `s.substituto` não alterou o Word
        público. Se um dia alguém declarar ali um substituto literal, este teste
        avisa, em vez de o arquivo mudar calado.
        """
        do_p4 = [s for s in S.carregar() if s.fonte == "produto4"]
        self.assertTrue(do_p4, "nenhuma supressão de produto4 no codebook")
        for s in do_p4:
            self.assertEqual(s.substituto, S.MARCA, f"§{s.paragrafo}")


class TestEquivalenciaEntrega(unittest.TestCase):
    """A trava do anexo de entrega: arquivo e página têm de dizer o mesmo.

    Se divergirem, quem lê o `.docx` e quem lê a página veem proteções
    diferentes para o mesmo trecho, que é o defeito que a ferramenta existe
    para fechar.
    """

    def comparar(self, esperado, obtido):
        from tools.transcricoes_entrega import comparar
        return comparar(esperado, obtido, "ENT-999")

    def test_iguais_passam(self):
        pars = ["cabeçalho", "", "um turno qualquer."]
        self.assertIsNone(self.comparar(pars, list(pars)))

    def test_texto_divergente_aborta_nomeando_o_indice(self):
        esperado = ["cabeçalho", "", "o trecho tratado (…) segue."]
        obtido = ["cabeçalho", "", "o trecho NÃO tratado segue."]
        with self.assertRaises(SystemExit) as caso:
            self.comparar(esperado, obtido)
        self.assertIn("§2", str(caso.exception))

    def test_contagem_divergente_aborta(self):
        with self.assertRaises(SystemExit):
            self.comparar(["a", "b"], ["a"])


class TestListaPessoas(unittest.TestCase):
    """A lista de participantes não pode absorver linha de cabeçalho.

    A aba PESSOAS da planilha real repete o cabeçalho no meio, onde começa um
    segundo bloco. Sem guarda, a palavra «Nome» entrava como se fosse gente, e a
    varredura reprovava qualquer documento com uma tabela de coluna «Nome».
    """

    def _planilha(self, caminho, linhas):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "PESSOAS"
        ws.append(["#", "Nome", "Email"])
        for i, n in enumerate(linhas, start=1):
            ws.append([i, n, ""])
        wb.save(caminho)

    def carregar(self, linhas, minimo=1):
        from tools.lista_pessoas import carregar
        with TemporaryDirectory() as td:
            p = os.path.join(td, "lista.xlsx")
            self._planilha(p, linhas)
            return carregar(p, minimo=minimo)

    def test_cabecalho_repetido_no_meio_nao_vira_pessoa(self):
        nomes = self.carregar(["Joana Fictícia", "Nome", "Pedro Fictício"])
        self.assertNotIn("Nome", nomes)
        self.assertIn("Joana Fictícia", nomes)
        self.assertIn("Pedro Fictício", nomes)

    def test_nomes_depois_do_cabecalho_repetido_continuam_sendo_lidos(self):
        """O segundo bloco é gente de verdade: descartar a linha, não o bloco."""
        nomes = self.carregar(["Joana Fictícia", "Nome", "Ana Fictícia",
                               "Carlos Fictício"])
        self.assertEqual(len(nomes), 3)

    def test_outros_rotulos_de_cabecalho_tambem_saem(self):
        for rotulo in ("Participante", "Entrevistado", "NOME COMPLETO"):
            with self.subTest(rotulo=rotulo):
                nomes = self.carregar(["Joana Fictícia", rotulo])
                self.assertNotIn(rotulo, nomes)


if __name__ == "__main__":
    unittest.main(verbosity=2)
