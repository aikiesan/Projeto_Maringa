# -*- coding: utf-8 -*-
"""Gera as 17 transcricoes em `.docx` com os cortes DENTRO do arquivo.

    python -m tools.transcricoes_entrega

Por que existe. Ate aqui as supressoes declaradas viviam em dois lugares: no
Anexo 05, para os sete cortes antigos, e na montagem da pagina, onde
`tools/hub/transcricoes.py` chama `aplicar_tolerante` na hora de converter. Quem
recebesse o `.docx` recebia o arquivo sem os cortes novos. O entregavel pedido e
a transcricao corrigida como ARQUIVO, e nao como pagina.

As 17 sao geradas, inclusive as que nao tem supressao declarada. O anexo de
entrega sao 17 sessoes; um diretorio com 13 arquivos seria ambiguo, e a ausencia
de um arquivo nao se distingue de um esquecimento.

A saida fica em `saida_entrega/`, fora do git, pela mesma regra que mantem
`saida_anonimizacao/` fora: e material de acervo.

## A trava

O texto do `.docx` gerado tem de ser, paragrafo a paragrafo, igual ao que a
pagina publica. Se divergirem, o leitor do arquivo e o leitor da pagina veem
protecoes diferentes, que e exatamente o defeito que esta ferramenta existe para
fechar. A conferencia roda sobre o arquivo salvo, nao sobre o objeto em memoria,
e interrompe na primeira divergencia nomeando o indice.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from tools.acervo import fonte_transcricoes, codigo, SUFIXO   # noqa: E402
from tools import supressoes as S                             # noqa: E402

DESTINO = ROOT / "saida_entrega" / "transcricoes"


def _textos(caminho) -> list[str]:
    import docx
    return [p.text for p in docx.Document(str(caminho)).paragraphs]


def _conferir(origem: Path, destino: Path, cod: str) -> None:
    """O arquivo salvo tem de dizer o mesmo que a pagina publicada.

    O esperado vem de `aplicar_tolerante` sobre os paragrafos CRUS da origem,
    que e literalmente o que `tools/hub/transcricoes.py` faz antes de montar o
    HTML. Comparar indice a indice, e nao o texto concatenado, porque e no
    indice cru que as declaracoes de `supressoes.csv` se ancoram.
    """
    esperado, _, _ = S.aplicar_tolerante(_textos(origem), "transcricao", cod)
    comparar(esperado, _textos(destino), cod)


def comparar(esperado: list[str], obtido: list[str], cod: str) -> None:
    """A comparacao, separada do disco para que o teste negativo a alcance.

    Uma trava sem teste negativo e uma trava que pode nao existir: ja aconteceu
    aqui, uma substituicao sem `assert` deixou uma guarda inteira por instalar e
    so o teste mostrou.
    """
    if len(esperado) != len(obtido):
        raise SystemExit(
            f"{cod}: o arquivo gerado tem {len(obtido)} paragrafos e a origem "
            f"{len(esperado)}. Aplicar supressao nao pode criar nem apagar "
            f"paragrafo. Nada foi entregue.")
    for i, (e, o) in enumerate(zip(esperado, obtido)):
        if e != o:
            raise SystemExit(
                f"{cod} §{i}: o .docx de entrega diverge da pagina publicada.\n"
                f"  pagina:  {e[:90]!r}\n"
                f"  arquivo: {o[:90]!r}\n"
                f"Duas protecoes diferentes para o mesmo trecho. Nada foi entregue.")


def gerar(destino: Path = DESTINO) -> dict:
    import docx
    destino.mkdir(parents=True, exist_ok=True)
    sups = S.carregar()
    linhas, tot_novas, tot_ja = [], 0, 0
    with fonte_transcricoes() as origem:
        if origem is None:
            raise SystemExit(
                "acervo indisponivel: nem saida_anonimizacao/ nem o Anexo 05. "
                "Gerar 0 arquivos e falhar, nao entregar vazio.")
        arquivos = sorted(origem.glob("*.docx"))
        if not arquivos:
            raise SystemExit(f"nenhum .docx em {origem}")
        for f in arquivos:
            cod = codigo(f)
            doc = docx.Document(str(f))
            novas, ja = S.aplicar_em_doc(doc.paragraphs, "transcricao", cod, sups)
            alvo = destino / (cod + SUFIXO)
            doc.save(str(alvo))
            _conferir(f, alvo, cod)
            linhas.append((cod, len(doc.paragraphs), novas, ja))
            tot_novas += novas
            tot_ja += ja
    return {"arquivos": len(linhas), "novas": tot_novas, "ja": tot_ja,
            "linhas": linhas, "destino": destino}


def main() -> int:
    r = gerar()
    print(f"transcricoes de entrega em {r['destino']}")
    for cod, pars, novas, ja in r["linhas"]:
        marca = f"  {novas} aplicada(s), {ja} ja estava(m)" if (novas or ja) else ""
        print(f"  {cod:16} {pars:5} paragrafos{marca}")
    print(f"\n  {r['arquivos']} arquivos · {r['novas']} supressoes aplicadas · "
          f"{r['ja']} ja presentes no anexo")
    print("  equivalencia com a pagina publicada conferida em todas as sessoes")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
