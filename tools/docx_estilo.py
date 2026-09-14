# -*- coding: utf-8 -*-
"""Estilo de marca para os `.docx` gerados do zero.

Existia um conjunto de helpers assim, mas aninhado dentro de `create_report()`
em `generate_report.py`, na raiz, com o texto do Produto 3 literal no codigo e o
nome do arquivo de saida fixo: nao dava para importar. Aqui eles ficam
importaveis, com a paleta do projeto, e com o helper de tabela que faltava (nao
havia um unico `add_table` no repositorio).

A paleta e a da identidade do projeto, e nao a das series dos graficos: cor de
marca e cor de dado sao coisas diferentes, e a das series nao se mexe sem medir
contraste.
"""
from __future__ import annotations

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

PETROLEO = RGBColor(0x0F, 0x4C, 0x5C)
VERDE = RGBColor(0x4E, 0x9F, 0x3D)
GRAFITE = RGBColor(0x4A, 0x4A, 0x4A)
TINTA = RGBColor(0x22, 0x22, 0x22)
MUDO = RGBColor(0x66, 0x66, 0x66)

FONTE = "Calibri"


def preparar(doc, margem_pol: float = 1.0) -> None:
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(margem_pol)
        s.left_margin = s.right_margin = Inches(margem_pol)
    normal = doc.styles["Normal"]
    normal.font.name = FONTE
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TINTA


def _fmt(run, tam=10.5, cor=TINTA, negrito=False, italico=False):
    run.font.name = FONTE
    run.font.size = Pt(tam)
    run.font.color.rgb = cor
    run.bold = negrito
    run.italic = italico
    return run


def titulo(doc, texto, sub=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _fmt(p.add_run(texto), 19, PETROLEO, negrito=True)
    p.paragraph_format.space_after = Pt(2)
    if sub:
        q = doc.add_paragraph()
        _fmt(q.add_run(sub), 11, GRAFITE)
        q.paragraph_format.space_after = Pt(14)
    return p


def h1(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    _fmt(p.add_run(texto), 14, PETROLEO, negrito=True)
    return p


def h2(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(3)
    _fmt(p.add_run(texto), 12, VERDE, negrito=True)
    return p


def corpo(doc, texto, prefixo_negrito="", mudo=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.15
    if prefixo_negrito:
        _fmt(p.add_run(prefixo_negrito), negrito=True)
    _fmt(p.add_run(texto), cor=MUDO if mudo else TINTA)
    return p


def item(doc, texto, prefixo_negrito=""):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    if prefixo_negrito:
        _fmt(p.add_run(prefixo_negrito), negrito=True)
    _fmt(p.add_run(texto))
    return p


def tabela(doc, cabecalho, linhas, larguras=None, num_a_direita=()):
    """Tabela com cabeçalho em negrito e colunas numéricas alinhadas à direita.

    `num_a_direita` são os índices das colunas numéricas. Alinhar número à
    direita não é enfeite: é o que deixa a coluna comparável na leitura.
    """
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for j, texto in enumerate(cabecalho):
        cel = t.rows[0].cells[j]
        cel.text = ""
        p = cel.paragraphs[0]
        p.alignment = (WD_ALIGN_PARAGRAPH.RIGHT if j in num_a_direita
                       else WD_ALIGN_PARAGRAPH.LEFT)
        _fmt(p.add_run(str(texto)), 9.5, PETROLEO, negrito=True)
    for linha in linhas:
        cells = t.add_row().cells
        for j, valor in enumerate(linha):
            cells[j].text = ""
            p = cells[j].paragraphs[0]
            p.alignment = (WD_ALIGN_PARAGRAPH.RIGHT if j in num_a_direita
                           else WD_ALIGN_PARAGRAPH.LEFT)
            _fmt(p.add_run("" if valor is None else str(valor)), 9.5)
    if larguras:
        for j, w in enumerate(larguras):
            for linha in t.rows:
                linha.cells[j].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def nota(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    _fmt(p.add_run(texto), 9, MUDO, italico=True)
    return p
