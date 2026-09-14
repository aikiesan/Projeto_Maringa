# -*- coding: utf-8 -*-
"""Produto 04 completo, editavel, sem travessoes. Versao NAO publica.

    python -m tools.p4_docx_completo

## Em que difere da versao publica

A versao publica (`tools/p4_docx_publico.py`) faz tres coisas ao texto: aplica as
supressoes declaradas, troca os quatro quadros nominais pela contagem que eles
sustentavam, e reescreve os travessoes. Esta faz **so a terceira**.

O arquivo aqui vai para IPPLAM e CEPAL pelo canal da entrega, junto da lista
nominal de entrevistados, que e entregavel previsto no termo de referencia. Para
quem ja tem a lista, suprimir o nome de um agente publico no corpo do relatorio
nao protege nada: so torna o documento mais dificil de conferir. O que a
supressao publica impede e a ligacao entre pessoa e sessao para quem NAO tem a
lista, e essa e uma preocupacao da camada publica.

Por isso o criterio e explicito: **este arquivo nao e menos protegido por
descuido, e sim por destino**. Ele nao pode ser publicado, e nao entra no git.

## O que muda em relacao a `anexos/Produto_4.docx`

So a pontuacao: 39 travessoes em 24 paragrafos, reescritos por virgula,
dois-pontos ou parenteses conforme o caso, preservando os runs, o que mantem o
negrito do rotulo que abre cada afirmacao da Secao 3. O diff foi conferido antes
de a passagem ser ligada, em 14/09.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

ORIGEM = ROOT / "anexos" / "Produto_4.docx"
DESTINO = ROOT / "saida_entrega" / "Produto_04_completo.docx"


def gerar(origem: Path = ORIGEM, destino: Path = DESTINO) -> dict:
    import docx
    sys.path.insert(0, str(ROOT / "tools" / "hub"))
    import produto4 as P4
    from tools.p4_travessoes import aplicar_no_doc

    doc = docx.Document(str(origem))
    # A mesma trava da versao publica: existe uma copia de 328 paragrafos que
    # termina em «3.4.5. Solucoes Propostas», sem o capitulo 4, as Consideracoes
    # Finais e os Anexos, e ela e MAIS NOVA no disco.
    P4.guarda(doc, origem)

    r = aplicar_no_doc(doc)
    quadros = sum(1 for t in doc.tables)

    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(destino))
    return {"travessoes": r["antes"], "pontuados": r["paragrafos"],
            "paragrafos": len(doc.paragraphs), "tabelas": quadros,
            "destino": destino, "bytes": destino.stat().st_size}


def main() -> int:
    r = gerar()
    print(f"{r['destino']}")
    for k in ("travessoes", "pontuados", "paragrafos", "tabelas", "bytes"):
        print(f"  {k:12} {r[k]}")
    print("  quadros nominais PRESERVADOS · supressoes NAO aplicadas")
    print("  FORA do git e fora do Hub. Entrega por canal proprio.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
