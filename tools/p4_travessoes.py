# -*- coding: utf-8 -*-
"""Reescreve os travessoes do Produto 4 na versao publica, preservando os runs.

    python -m tools.p4_travessoes            # mostra o diff, nao escreve nada

Por que existe. O pedido de 13/09 e que a escrita do projeto nao tenha cara de
gerada por IA, e o travessao e a marca mais visivel disso. As paginas do Hub ja
saem sem nenhum; o texto do relatorio tem 39, em 24 paragrafos, e eles vem do
proprio `.docx`. Retira-los e editar o texto do autor, e por isso esta
ferramenta mostra o diff e so o gerador do arquivo publico aplica.

## As tres regras, e por que sao tres

1. **Titulo de anexo.** «Anexo 01 — Instrumentais» vira dois-pontos. Sao seis, e
   e o unico caso puramente tipografico.
2. **Aposto entre dois travessoes.** Vira parenteses, que e a pontuacao que faz
   o mesmo trabalho sem a marca. Quando o trecho entre travessoes JA tem
   parenteses dentro, aninhar seria pior do que o problema, e entao a troca e
   por virgulas.
3. **Travessao isolado.** Vira virgula quando o que vem depois e coordenacao
   («e», «mas», «ou»), porque virgula e o que liga oracoes coordenadas; vira
   dois-pontos nos demais, porque o que vem depois e explicacao ou enumeracao, e
   e isso que dois-pontos anuncia.

## A trava

Sobrar um travessao e falha: a ferramenta conta antes e depois e levanta se a
conta nao fechar, em vez de entregar um arquivo meio tratado.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

TRAVESSAO = "—"
ORIGEM = ROOT / "anexos" / "Produto_4.docx"

# O que separa um aposto de dois travessoes isolados nao e o tamanho: e o ponto
# final. Medido nos 39: os apostos tem 43, 48, 64, 108, 126 e 180 caracteres e
# nenhum ponto final dentro; os nao-apostos tem 105, 189 e 513 e TODOS tem. Em
# §333 o primeiro travessao abre uma ressalva que atravessa o periodo inteiro.
# O limite de tamanho fica so como rede contra caso patologico.
MAX_APOSTO = 250

_ANEXO = re.compile(r"(Anexo \d+) " + TRAVESSAO + " ")
_PAR = re.compile(" " + TRAVESSAO + r" ([^" + TRAVESSAO + r"]+?) " + TRAVESSAO
                  + r"(?=[\s,.;:])")
_PAR_VIRG = re.compile(" " + TRAVESSAO + r" ([^" + TRAVESSAO + r"]+?) "
                       + TRAVESSAO + r"(?=,)")
_COORD = re.compile(r"^(e|mas|ou|nem|porem|porém)$", re.IGNORECASE)


def _aposto(m: re.Match) -> str:
    dentro = m.group(1)
    if "(" in dentro or ")" in dentro:
        # parenteses dentro de parenteses e pior do que o travessao
        return ", " + dentro + ","
    return " (" + dentro + ")"


def _isolado(m: re.Match) -> str:
    """Virgula quando o que vem depois e coordenacao; dois-pontos nos demais.

    A primeira versao capturava `(\S)`, um unico caractere, e entao «mas» chegava
    aqui como «m»: nenhuma coordenacao era reconhecida e tudo virava dois-pontos.
    Capturar a palavra inteira e o que faz a regra existir.
    """
    palavra = m.group(1)
    return (", " if _COORD.match(palavra) else ": ") + palavra


def reescrever(texto: str) -> str:
    """Aplica as tres regras a um paragrafo. Idempotente."""
    if TRAVESSAO not in texto:
        return texto
    t = _ANEXO.sub(r"\1: ", texto)

    def _par_sub(m):
        dentro = m.group(1)
        if len(dentro) > MAX_APOSTO or ". " in dentro:
            return m.group(0)          # nao e aposto; deixa para a regra 3
        return _aposto(m)

    def _par_virg_sub(m):
        """O par que fecha antes de virgula («… se aplica —, enquanto»).

        A virgula ja esta no texto, entao a troca por virgulas nao repoe a do
        fim: sairia «equivalentes),,». Esta funcao existia sem o teste de
        parenteses internos, e §330 saia com parentese dentro de parentese.
        """
        dentro = m.group(1)
        if len(dentro) > MAX_APOSTO or ". " in dentro:
            return m.group(0)
        if "(" in dentro or ")" in dentro:
            return ", " + dentro
        return " (" + dentro + ")"

    t = _PAR_VIRG.sub(_par_virg_sub, t)
    t = _PAR.sub(_par_sub, t)
    t = re.sub(" " + TRAVESSAO + r" (\S+)", _isolado, t)
    return t


def paragrafos_afetados(doc):
    return [(i, p) for i, p in enumerate(doc.paragraphs) if TRAVESSAO in p.text]


def aplicar_no_doc(doc) -> dict:
    """Reescreve no documento, preservando os runs. Levanta se sobrar travessao."""
    antes = sum(p.text.count(TRAVESSAO) for p in doc.paragraphs)
    n_par = 0
    for _, p in paragrafos_afetados(doc):
        novo = reescrever(p.text)
        if novo == p.text:
            continue
        _escrever_preservando_runs(p, novo)
        n_par += 1
    depois = sum(p.text.count(TRAVESSAO) for p in doc.paragraphs)
    if depois:
        restantes = [f"§{i}" for i, p in paragrafos_afetados(doc)]
        raise SystemExit(
            f"sobraram {depois} travessao(oes) em {', '.join(restantes)}. "
            f"Entregar o texto meio tratado e pior do que nao tratar.")
    return {"antes": antes, "depois": depois, "paragrafos": n_par}


def _escrever_preservando_runs(par, novo: str) -> None:
    """Poe `novo` no paragrafo mexendo no MENOR numero de runs possivel.

    O negrito que abre cada afirmacao da Secao 3 vive num run separado. Reescrever
    o paragrafo inteiro o apagaria. Como a reescrita so troca pontuacao, o prefixo
    e o sufixo comuns sao longos: basta achar o trecho que de fato mudou e
    reescrever os runs que ele atravessa, como faz `supressoes.aplicar_em_paragrafo`.
    """
    velho = par.text
    i = 0
    while i < min(len(velho), len(novo)) and velho[i] == novo[i]:
        i += 1
    j, k = len(velho), len(novo)
    while j > i and k > i and velho[j - 1] == novo[k - 1]:
        j -= 1
        k -= 1
    meio = novo[i:k]
    if not par.runs:
        par.add_run(novo)
        return
    pos, primeiro = 0, True
    for r in par.runs:
        a, b = pos, pos + len(r.text)
        pos = b
        if b <= i or a >= j:
            continue
        ini = r.text[:max(0, i - a)]
        fim = r.text[max(0, j - a):] if j - a < len(r.text) else ""
        r.text = ini + (meio if primeiro else "") + fim
        primeiro = False
    if primeiro:                        # nada foi atravessado: caso defensivo
        par.runs[-1].text += meio


def main() -> int:
    import docx
    doc = docx.Document(str(ORIGEM))
    afetados = paragrafos_afetados(doc)
    total = sum(p.text.count(TRAVESSAO) for _, p in afetados)
    print(f"{total} travessoes em {len(afetados)} paragrafos de {ORIGEM.name}\n")
    mudou = 0
    for i, p in afetados:
        novo = reescrever(p.text)
        if novo == p.text:
            print(f"§{i}  NAO TRATADO: {p.text[:110]}")
            continue
        mudou += 1
        print(f"§{i}")
        for a, b in _trechos_que_mudaram(p.text, novo):
            print(f"   - {a}")
            print(f"   + {b}")
        print()
    resta = sum(reescrever(p.text).count(TRAVESSAO) for _, p in afetados)
    print(f"{mudou} paragrafos reescritos · {resta} travessoes restantes")
    return 1 if resta else 0


def _trechos_que_mudaram(velho: str, novo: str, ctx: int = 42):
    """Só os pedaços que mudaram, com contexto. O parágrafo inteiro é ilegível.

    Alinhamento por `difflib`, e não por busca de prefixo: a primeira versão
    procurava o contexto antigo dentro do texto novo e imprimia «(?)» sempre que
    o próprio contexto tivesse sido reescrito, que é justamente o caso do
    segundo travessão de um par.
    """
    import difflib
    saida = []
    sm = difflib.SequenceMatcher(None, velho, novo, autojunk=False)
    for tag, a1, a2, b1, b2 in sm.get_opcodes():
        if tag == "equal":
            continue
        saida.append((velho[max(0, a1 - ctx):a2 + ctx].strip(),
                      novo[max(0, b1 - ctx):b2 + ctx].strip()))
    return saida


if __name__ == "__main__":
    raise SystemExit(main())
