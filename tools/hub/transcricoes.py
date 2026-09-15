# -*- coding: utf-8 -*-
"""Converte as transcrições anonimizadas (.docx) em páginas do Hub."""
import re, os, glob, html
from pathlib import Path
from docx import Document

E = html.escape
ROTULO_SETOR = {"Publico": "Público"}
TS = re.compile(r"^\d{2}:\d{2}:\d{2}$")
FALA = re.compile(r"^(\[(?:entrevistador|participante)[^\]]*\]|\[nome\])\s*:\s*(.*)$", re.S)

CSS_EXTRA = """
<style>
.transc{max-width:none}
.turno{display:grid; grid-template-columns:132px 1fr; gap:14px; padding:7px 0;
  border-bottom:1px solid var(--grid); align-items:start}
.turno:last-child{border-bottom:0}
.turno .quem{font-size:12.5px; font-weight:600; color:var(--ink2); padding-top:2px}
.turno .quem.e{color:var(--s2)}
.turno .txt{max-width:72ch}
.marca-t{font-variant-numeric:tabular-nums; font-size:12px; color:var(--mut);
  margin:22px 0 6px; letter-spacing:.02em}
.corte{color:var(--mut); font-style:italic}
@media (max-width:640px){
  .turno{grid-template-columns:1fr; gap:2px}
  .turno .quem{padding-top:0}
}
</style>"""


def converte(caminho):
    """Devolve (codigo, cabecalho_dict, html_do_corpo)."""
    cod = os.path.basename(caminho).replace("_Transcricao_anonimizada.docx", "")
    doc = Document(caminho)

    # As supressoes declaradas sao aplicadas AQUI, na geracao da pagina, e nao
    # so no pipeline de anonimizacao. A razao: o Anexo 05 que esta ferramenta le
    # e um zip ja gerado, e uma declaracao nova em codebook/supressoes.csv so
    # chegaria ao publico depois de regerar o anexo inteiro, o que depende do
    # acervo e do vault. Aplicando na camada publicada, a pagina respeita a
    # declaracao mesmo com o anexo velho.
    #
    # Os indices das declaracoes sao de `doc.paragraphs` cru. O filtro de
    # paragrafos vazios vem DEPOIS, senao os indices deslizam e a supressao cai
    # no paragrafo errado.
    brutos = [p.text for p in doc.paragraphs]
    import sys as _sys
    _sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
    from tools.supressoes import aplicar_tolerante
    brutos, _novas, _ja = aplicar_tolerante(brutos, "transcricao", cod)
    if _novas:
        print(f"    {cod}: {_novas} supressao(oes) aplicada(s) na publicacao"
              + (f", {_ja} ja vinham do anexo" if _ja else ""))

    paras = [t.strip() for t in brutos if t.strip()]

    cab = {}
    if paras and paras[0].startswith("TRANSCRIÇÃO ANONIMIZADA"):
        linhas = paras[0].split("\n")
        for ln in linhas:
            for campo in ("Setor", "Data", "Duração", "Participantes"):
                m = re.search(campo + r":\s*([^·]+)", ln)
                if m:
                    cab[campo] = m.group(1).strip()
        paras = paras[1:]

    out, n_turnos = [], 0
    for t in paras:
        if TS.match(t):
            # ancora estavel para o deep-link vindo do Produto 4:
            # transcricoes/ENT-XXX.html#t-002243
            out.append(f'<p class="marca-t" id="t-{t.replace(":", "")}">{t}</p>')
            continue
        if t.startswith("Entrevista ") or t.startswith("Esta transcrição") \
           or re.match(r"^[a-z]{3}\.? \d{1,2}, \d{4}$", t):
            continue
        m = FALA.match(t)
        if m:
            quem, fala = m.group(1), m.group(2)
            cls = "quem e" if "entrevistador" in quem else "quem"
            fala = fala.replace("(…)", '<span class="corte">(…)</span>') \
                if "(…)" in fala else E(fala)
            if "(…)" in t:
                fala = E(m.group(2)).replace("(…)", '<span class="corte">(…)</span>')
            n_turnos += 1
            out.append(f'<div class="turno"><span class="{cls}">{E(quem)}</span>'
                       f'<span class="txt">{fala}</span></div>')
        else:
            out.append(f'<div class="turno"><span class="quem"></span>'
                       f'<span class="txt">{E(t)}</span></div>')
    cab["turnos"] = n_turnos
    return cod, cab, "".join(out)


def pagina_transcricao(cod, cab, corpo, sessao=None):
    """O cabeçalho vem da camada SANEADA (dados.sessoes), não do .docx — que foi
    escrito antes do saneamento e ainda traz o setor «Especial» e o rótulo
    institucional preciso."""
    if sessao:
        meta = (f"Setor: {ROTULO_SETOR.get(sessao['setor_publico'], sessao['setor_publico'])} · "
                ""    # tipo institucional removido: ver tools/hub/dados.py
                f"Data: {sessao['date']} · Duração: {sessao['minutes']} min · "
                f"Participantes: {sessao['n_participants']}")
    else:
        meta = " · ".join(f"{k}: {v}" for k, v in cab.items()
                          if k in ("Setor", "Data", "Duração"))
    return f"""
<p class="sm"><a href="transcricoes.html">&larr; Todas as transcrições</a></p>
<h1>{cod}</h1>
<p class="lede">{E(meta)}</p>

<div class="nota">
  <p><strong>Camada anonimizada.</strong> Nomes de pessoas foram substituídos por
  <code>[entrevistador n]</code>, <code>[participante n]</code> e <code>[nome]</code>;
  endereços de e-mail, telefones, identificadores fiscais e links foram suprimidos.
  Órgãos públicos, leis, instrumentos de política e contratos publicados permanecem
  nominais, por serem informação pública. As marcas de tempo são as do registro
  original e ancoram as evidências codificadas. Reticências entre parênteses
  <span class="corte">(…)</span> indicam trecho suprimido por risco de
  reidentificação. O registro integral dos cortes é mantido pela
  consultoria.</p>
</div>

<div class="transc">{corpo}</div>
"""


# A marca das sessões retidas é genérica e igual para todas: o porquê mora em
# `codebook/retencao_publica.csv`, que não é publicado. Dizer aqui o motivo, o
# critério ou o número seria repor por outro caminho o rótulo por sessão que
# saiu da camada pública em 13/09.
MARCA_RETIDA = "não publicada nesta camada · consta do anexo de entrega"


def indice(sessoes, cabecalhos, publicadas=None):
    """O índice lista as 17, com as retidas marcadas e sem link.

    Listar só as publicadas contradiria o resto do Hub, que conta 17 sessões em
    todos os números, e deixaria sem resposta quem seguisse uma evidência do
    painel até a transcrição.
    """
    linhas, n_retidas = [], 0
    for s in sorted(sessoes, key=lambda x: x["code"]):
        cod = s["code"]
        cab = cabecalhos.get(cod, {})
        setor = ROTULO_SETOR.get(s["setor_publico"], s["setor_publico"])
        publica = publicadas is None or cod in publicadas
        if publica:
            celula = f'<td><a href="transcricoes/{cod}.html"><strong>{cod}</strong></a></td>'
        else:
            n_retidas += 1
            celula = (f'<td><strong>{cod}</strong>'
                      f'<br><span class="sm mut">{MARCA_RETIDA}</span></td>')
        linhas.append(
            f'<tr data-g="{E(setor)}">'
            f'{celula}'
            f'<td>{E(setor)}</td>'
            f'<td class="num">{E(s["date"])}</td>'
            f'<td class="num">{E(s["minutes"])} min</td>'
            f'<td class="num">{cab.get("turnos", "n/d")}</td>'
            f'<td class="num">{E(s["n_participants"])}</td></tr>')
    # Os dois números da nota são derivados, não digitados: um vem da contagem
    # de retidas feita acima, o outro do codebook de supressões.
    from tools.supressoes import carregar as _sups
    # Conta só entre as sessões que o leitor tem diante de si: falar dos cortes
    # das retidas seria descrever texto que não está nesta página.
    n_com_corte = len({x.alvo for x in _sups() if x.fonte == "transcricao"
                       and (publicadas is None or x.alvo in publicadas)})
    retidas_frase = (f"; {n_retidas} integram o anexo de entrega e não esta camada"
                     if n_retidas else "")
    setores = sorted({ROTULO_SETOR.get(s["setor_publico"], s["setor_publico"])
                      for s in sessoes})
    opts = "".join(f'<option value="{E(g)}">{E(g)}</option>' for g in setores)
    tot_min = sum(int(s["minutes"]) for s in sessoes)
    return f"""
<h1>Transcrições</h1>

<p class="lede">As {len(sessoes)} sessões de entrevista, em camada anonimizada,
somando {tot_min // 60}h{tot_min % 60:02d} de escuta institucional com
{sum(int(s['n_participants']) for s in sessoes)} participantes dos quatro grupos de
atores-chave. {len(sessoes) - n_retidas} estão publicadas na íntegra
aqui{retidas_frase}.</p>

<div class="nota">
  <p><strong>O que foi feito com estes textos.</strong> Nomes de participantes, de
  entrevistadores e de terceiros citados foram substituídos por rótulos; contatos e
  links, suprimidos. Em {n_com_corte} sessões há trechos suprimidos porque o cargo, a
  filiação declarada ou o local de trabalho identificavam o participante
  independentemente do nome; cada corte está marcado no texto e registrado em relação
  anexa ao produto.</p>
  <p><strong>Por que {n_retidas} sessões não estão aqui.</strong> Medimos quantas vezes
  o participante nomeia o próprio empregador junto de uma marca de primeira pessoa.
  Nas sessões em que isso é recorrente, publicar a transcrição integral e proteger o
  local de trabalho de quem falou por uma hora sobre o próprio trabalho são objetivos
  incompatíveis: seriam necessários tantos cortes que o conteúdo não sobreviveria.
  Essas sessões foram retiradas desta camada e integram o anexo de entrega à
  coordenação, com o critério e a decisão registrados no codebook do projeto.</p>
  <p>Subsiste risco residual de reidentificação por parte de quem conheça a estrutura
  institucional do município: ele é assumido e decorre da própria natureza de um
  corpus em que a representatividade setorial exige ouvir quem ocupa posição única
  na estrutura municipal.</p>
</div>

<div class="busca">
  <input id="q" type="search" placeholder="Buscar sessão…" aria-label="Buscar">
  <select id="fg" aria-label="Filtrar por setor">
    <option value="">Todos os setores</option>{opts}
  </select>
</div>
<p class="conta" id="conta"></p>
<div class="tw"><table id="tab">
<thead><tr><th>Sessão</th><th>Setor</th>
<th class="num">Data</th><th class="num">Duração</th><th class="num">Turnos</th>
<th class="num">Particip.</th></tr></thead>
<tbody>{"".join(linhas)}</tbody></table></div>

<p class="sm mut">Todas as sessões foram gravadas mediante Termo de Consentimento e
Autorização para Gravação de Áudio e Vídeo assinado pelo participante.</p>
"""
