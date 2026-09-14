# -*- coding: utf-8 -*-
"""Relatorio de codificacao dos resultados, em `.docx` editavel.

    python -m tools.relatorio_codificacao

O que e. O Produto 04 traz o diagnostico; este documento traz a CODIFICACAO que
o sustenta: quantas evidencias, de onde, como foram graduadas, o que ficou sem
evidencia e o que ficou divergente. E o documento que permite contestar uma
afirmacao do diagnostico sem ter de reabrir o codebook.

Tudo e derivado de `tools/p4_base.py` e `tools/hub/dados.py`, que leem os CSVs
versionados. Nenhum numero e digitado aqui: corrigir um numero significa
corrigir o codebook, nunca a saida.

## O que NAO pode entrar

Este relatorio descreve o corpus, e nao as pessoas. Nao traz rotulo
institucional de sessao, nem nada que ligue pessoa a sessao. A sessao aparece
so pelo codigo e pelo setor. Rode `tools/hub/varre_hub.py` sobre ele antes de
entregar: a varredura ja le `.docx`.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

DESTINO = ROOT / "hub_saida" / "Relatorio_Codificacao.docx"

# O codebook grava «Publico» sem acento; a camada de leitura acentua. Mesma
# tabela que `tools/hub/transcricoes.py` usa, pela mesma razao.
ROTULO_SETOR = {"Publico": "Público"}


def _setor(nome: str) -> str:
    return ROTULO_SETOR.get(nome, nome)


def _prioridades(cob):
    for chave, rotulo in (("prioridade_muito_alta", "Muito alta"),
                          ("prioridade_alta", "Alta"),
                          ("prioridade_média", "Média"),
                          ("prioridade_baixa", "Baixa")):
        d = cob.get(chave)
        if d:
            yield rotulo, d


def gerar(destino: Path = DESTINO) -> dict:
    import docx
    from tools import docx_estilo as S
    from tools import p4_base as B
    sys.path.insert(0, str(ROOT / "tools" / "hub"))
    import dados as D

    b = B.base()
    con, cob, subs, set_ = b["consultas"], b["cobertura"], b["subcategorias"], b["setores"]
    m = D.metricas_corpus()

    doc = docx.Document()
    S.preparar(doc)

    S.titulo(doc, "Codificação dos resultados",
             "Consultoria técnica IPPLAM, CEPAL/ONU · condições habilitantes ao "
             "financiamento climático urbano de Maringá · anexo ao Produto 04")
    S.corpo(doc,
            "Este documento descreve como o corpus de entrevistas foi codificado e o "
            "que a codificação produziu. Todos os números saem de contagem sobre os "
            "arquivos versionados do codebook e são reproduzíveis por "
            "python -m tools.p4_base. Nenhum valor foi digitado.")

    # ------------------------------------------------------------ base
    S.h1(doc, "1. Base da avaliação")
    S.tabela(doc,
             ("Indicador", "Valor"),
             [("Sessões de entrevista", con["sessoes"]),
              ("Participantes", con["atores"]),
              ("Duração total", f"{con['duracao']} ({con['minutos']} min)"),
              ("Turnos de fala", f"{con['turnos']:,}".replace(",", ".")),
              ("Evidências codificadas", con["evidencias"]),
              ("Dimensões da metodologia", cob["dimensoes_total"]),
              ("Dimensões com evidência", cob["dimensoes_com_evidencia"]),
              ("Dimensões trianguladas", cob["trianguladas"]),
              ("Dimensões de fonte única", cob["fonte_unica"]),
              ("Evidências divergentes", cob["divergencias"]),
              ("Subcategorias avaliadas", len(subs))],
             larguras=(3.6, 1.8), num_a_direita=(1,))
    S.corpo(doc, "Distribuição setorial das sessões: " + " · ".join(
        f"{_setor(k)} {v} sessões e {set_.get(k, {}).get('n', 0)} evidências"
        for k, v in sorted(con["por_setor"].items())) + ".")

    S.h2(doc, "1.1. Unidade de registro")
    S.corpo(doc,
            "A unidade de registro é a evidência, não a entrevista nem a resposta. "
            "Cada evidência é um trecho anonimizado, com marca de tempo, vinculado a "
            "uma única dimensão, com tipo, maturidade, alinhamento e confiança. A "
            "paráfrase é o juízo do codificador e é o que se cita; o trecho permanece "
            "como lastro verificável.")

    # ------------------------------------------------------ cobertura
    S.h1(doc, "2. Cobertura por prioridade")
    S.corpo(doc,
            "A prioridade é a atribuída a Maringá no Produto 02. A leitura relevante "
            "não é quantas dimensões foram cobertas, e sim se as de maior prioridade "
            "foram.")
    S.tabela(doc,
             ("Prioridade", "Dimensões", "Com evidência", "Evidências"),
             [(rot, d["dimensoes"], d["com_evidencia"], d["evidencias"])
              for rot, d in _prioridades(cob)],
             larguras=(1.8, 1.3, 1.5, 1.3), num_a_direita=(1, 2, 3))

    S.h2(doc, "2.1. Graduação de maturidade")
    S.tabela(doc, ("Estágio", "Evidências"),
             [(rot, n) for rot, n in m["maturidade"]],
             larguras=(2.8, 1.4), num_a_direita=(1,))
    fora = {k: v for k, v in m["maturidade_completa"]
            if k in ("sem evidência", "não aplicável")} or dict(
        x for x in m["maturidade_completa"] if x[0] not in dict(m["maturidade"]))
    S.nota(doc,
           "Os estágios somam " + str(sum(n for _, n in m["maturidade"]))
           + " e não " + str(con["evidencias"]) + ": "
           + " e ".join(f"{n} classificada{'s' if n > 1 else ''} como «{k}»"
                        for k, n in fora.items())
           + " ficam fora da escala. «Sem evidência» é pendência de reinquirição, "
             "não ausência constatada, e somá-la a «inexistente» transformaria "
             "pendência em achado.")

    # --------------------------------------------------------- matriz
    S.h1(doc, "3. A matriz das subcategorias")
    S.corpo(doc,
            "A maturidade da subcategoria é a moda das evidências graduadas, e não a "
            "média. A escala é ordinal: a média entre «inexistente» e «efetivo» não "
            "descreve nenhum estado real. Em empate, prevalece o estágio menor, que é "
            "a leitura conservadora. O perfil diz se a distribuição é concentrada, "
            "dispersa ou bimodal, e a bimodalidade é informação: significa que a "
            "mesma subcategoria tem instrumento formalizado e prática ausente.")
    S.tabela(doc,
             ("#", "Subcategoria", "Dim.", "Evid.", "Sessões", "Maturidade",
              "% inexist.", "Perfil"),
             [(s["subcat"], s["nome"], f"{s['n_cobertas']}/{s['n_dimensoes']}",
               s["n_evidencias"], s["n_sessoes"], s["maturidade_rotulo"],
               f"{s['share_inexistente']}%",
               (s["perfil"] + (": " + " / ".join(s["perfil_polos"])
                               if s.get("perfil_polos") else "")))
              for s in subs],
             larguras=(0.5, 2.1, 0.6, 0.6, 0.7, 1.2, 0.8, 1.5),
             num_a_direita=(2, 3, 4, 6))

    # ------------------------------------------------------ por eixo
    S.h1(doc, "4. Leitura por eixo")
    for eixo in sorted({s["eixo"] for s in subs}):
        doeixo = [s for s in subs if s["eixo"] == eixo]
        nome = doeixo[0]["eixo_nome"]
        ev = sum(s["n_evidencias"] for s in doeixo)
        cobertas = sum(s["n_cobertas"] for s in doeixo)
        totd = sum(s["n_dimensoes"] for s in doeixo)
        modas = [s["maturidade_rotulo"] for s in doeixo]
        dominante = max(set(modas), key=modas.count)
        S.corpo(doc,
                f"{len(doeixo)} subcategorias, {cobertas} de {totd} dimensões com "
                f"evidência, {ev} evidências. Maturidade modal predominante: "
                f"{dominante.lower()}. "
                + ("Há bimodalidade em "
                   + ", ".join(s["subcat"] for s in doeixo if s["perfil"] == "bimodal")
                   + "." if any(s["perfil"] == "bimodal" for s in doeixo)
                   else "Sem bimodalidade."),
                prefixo_negrito=f"Eixo {eixo}, {nome}. ")

    # ---------------------------------------------------- por setor
    S.h1(doc, "5. O que cada grupo de atores enxerga")
    S.corpo(doc,
            "A mesma dimensão recebe graduação diferente conforme quem fala. A "
            "divergência não é ruído a resolver: é o achado que a etapa de validação "
            "participativa tem de tratar.")
    S.tabela(doc,
             ("Grupo", "Evidências", "% inexistente", "Maturidade modal"),
             [(_setor(k), v["n"], f"{v['share_inexistente']}%", v["rotulo"])
              for k, v in sorted(set_.items())],
             larguras=(1.9, 1.3, 1.5, 1.7), num_a_direita=(1, 2))

    # ------------------------------------------------------- lacunas
    S.h1(doc, "6. Dimensões sem evidência")
    sem = cob["sem_evidencia"]
    S.corpo(doc,
            f"{len(sem)} das {cob['dimensoes_total']} dimensões atravessaram o corpus "
            "sem qualquer evidência codificada. Ausência de evidência é resultado, e "
            "não falha de coleta: uma dimensão de prioridade alta sobre a qual "
            "ninguém falou é achado do diagnóstico.")
    S.tabela(doc,
             ("Dimensão", "Prioridade", "Subcategoria", "Nome"),
             [(d["code"], d.get("prioridade", ""), d.get("subcat", ""), d.get("nome", ""))
              for d in sem],
             larguras=(0.9, 1.1, 0.9, 3.3))

    # ---------------------------------------------------- integridade
    S.h1(doc, "7. Integridade da codificação")
    S.item(doc, f"{cob['trianguladas']} dimensões têm evidência de mais de uma sessão; "
                f"{cob['fonte_unica']} têm fonte única e são assinaladas como tal.",
           prefixo_negrito="Triangulação. ")
    S.item(doc, f"{cob['divergencias']} evidências divergem entre si e foram "
                "preservadas, nunca resolvidas por predominância de fonte.",
           prefixo_negrito="Divergência. ")
    S.item(doc, f"{con['com_tcle']} das {con['sessoes']} sessões têm Termo de "
                "Consentimento arquivado."
                + (f" Pendentes: {', '.join(con['sem_tcle'])}." if con["sem_tcle"] else ""),
           prefixo_negrito="Consentimento. ")
    S.item(doc, "Nomes de participantes, de entrevistadores e de terceiros citados "
                "foram substituídos por rótulos; trechos autoidentificadores foram "
                "suprimidos e cada corte está declarado no codebook, com o contexto e "
                "o hash do trecho, nunca com o trecho.",
           prefixo_negrito="Anonimização. ")

    S.nota(doc, "Documento derivado do codebook versionado. Reproduzível por "
                "python -m tools.relatorio_codificacao.")

    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(destino))
    return {"paragrafos": len(doc.paragraphs), "tabelas": len(doc.tables),
            "bytes": destino.stat().st_size, "destino": destino}


def main() -> int:
    r = gerar()
    print(f"gerado: {r['destino']}")
    for k in ("paragrafos", "tabelas", "bytes"):
        print(f"  {k:12} {r[k]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
