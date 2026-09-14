# -*- coding: utf-8 -*-
"""Carga da lista de participantes, para as varreduras de publicação.

Existe uma só razão para este módulo: há DUAS cópias da `Lista de Entrevistas.xlsx`
no acervo, com nomes de aba diferentes — a de `DRIVE_FILES/` tem `PESSOAS` (51
nomes) e a da raiz tem `Planilha1` (25 nomes). Um `wb["PESSOAS"]` fixo levanta
`KeyError` na segunda, e um `except` genérico transforma isso num aviso: a
varredura segue com a lista vazia e sai 0. Uma varredura que não carregou a
lista não está limpa — está cega.

Por isso: a aba e a coluna são localizadas pelo CABEÇALHO, não pelo nome, e a
falha é dura. Quem varre chama `carregar()` e deixa a exceção subir.
"""
from __future__ import annotations

from pathlib import Path


class ListaIndisponivel(RuntimeError):
    """A lista de participantes não pôde ser carregada. Nunca ignore."""


def carregar(caminho: Path, minimo: int = 20) -> list[str]:
    """Devolve os nomes completos da lista. Levanta ListaIndisponivel se algo
    não bater — incluindo um total implausivelmente baixo, que denuncia aba
    errada ou planilha truncada."""
    caminho = Path(caminho)
    if not caminho.exists():
        raise ListaIndisponivel(f"planilha não encontrada: {caminho}")
    try:
        import openpyxl
    except ImportError as e:
        raise ListaIndisponivel(f"openpyxl ausente: {e}") from e

    try:
        wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    except Exception as e:
        raise ListaIndisponivel(f"não foi possível abrir {caminho.name}: {e}") from e

    nomes: set[str] = set()
    abas_lidas = []
    for ws in wb.worksheets:
        col = _coluna_nome(ws)
        if col is None:
            continue
        abas_lidas.append(ws.title)
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or col >= len(row):
                continue
            v = row[col]
            if isinstance(v, str):
                limpo = _so_o_nome(v)
                # A aba PESSOAS de «Lista de Entrevistas.xlsx» repete o
                # cabecalho na linha 39, onde comeca um segundo bloco. Sem esta
                # guarda, a palavra «Nome» entrava na lista como se fosse uma
                # pessoa, e a varredura reprovava qualquer documento com uma
                # tabela que tivesse coluna «Nome». Os nomes do segundo bloco
                # continuam sendo lidos: o que se descarta e so o cabecalho.
                if limpo and limpo.strip().lower() in _ROTULOS_CABECALHO:
                    continue
                if limpo:
                    nomes.add(limpo)

    # `read_only=True` mantem o arquivo aberto ate o close explicito. Sem isto o
    # handle fica pendurado enquanto o processo viver, e no Windows a planilha
    # nao pode ser movida nem apagada por mais ninguem.
    wb.close()

    if not abas_lidas:
        raise ListaIndisponivel(
            f"{caminho.name}: nenhuma aba com coluna de nome "
            f"(abas vistas: {wb.sheetnames})")
    if len(nomes) < minimo:
        raise ListaIndisponivel(
            f"{caminho.name}: só {len(nomes)} nomes nas abas {abas_lidas} — "
            f"esperado ao menos {minimo}. Aba errada ou planilha truncada?")
    return sorted(nomes)


_ROTULOS_CABECALHO = ("nome", "nome completo", "entrevistado", "participante")


def _coluna_nome(ws) -> int | None:
    """Índice 0-based da coluna cujo cabeçalho é «Nome», na primeira linha."""
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        for i, c in enumerate(row or ()):
            if isinstance(c, str) and c.strip().lower() in _ROTULOS_CABECALHO:
                return i
    return None


# caminhos conhecidos, relativos à raiz do projeto. As duas cópias divergem
# (51 e 25 nomes); para varredura, a união é estritamente mais segura — procurar
# um nome a mais nunca torna a saída menos protegida.
CAMINHOS = (
    "DRIVE_FILES/drive-download-20260906T102534Z-1-001/Lista de Entrevistas.xlsx",
    "Lista de Entrevistas.xlsx",
)


def carregar_todas(raiz: Path, caminhos=CAMINHOS) -> list[str]:
    """União dos nomes de todas as cópias encontradas. Levanta
    ListaIndisponivel se NENHUMA carregar — é o caso em que varrer não faz
    sentido e o chamador tem de reprovar."""
    raiz = Path(raiz)
    nomes: set[str] = set()
    erros = []
    for rel in caminhos:
        try:
            nomes.update(carregar(raiz / rel))
        except ListaIndisponivel as e:
            erros.append(str(e))
    if not nomes:
        raise ListaIndisponivel(
            "nenhuma cópia da lista de participantes pôde ser carregada:\n  - "
            + "\n  - ".join(erros))
    return sorted(nomes)


# «Alef (Defesa Civil)» é nome + anotação de instituição. O código anterior
# trocava os parênteses por espaço e promovia «Defesa» e «Civil» a token de
# nome — que é exatamente o ruído por token isolado que já custou tempo neste
# projeto. O parêntese é descartado: instituição não é nome de pessoa.
import re as _re

_PARENTESE = _re.compile(r"\([^)]*\)")


def _so_o_nome(valor: str) -> str:
    limpo = _PARENTESE.sub(" ", valor)
    limpo = _re.sub(r"\s+", " ", limpo).strip(" -,;")
    return limpo
