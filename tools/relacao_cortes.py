# -*- coding: utf-8 -*-
"""Relacao de cortes por identificabilidade, em planilha, para a coordenacao.

    python -m tools.relacao_cortes

## Por que existe

As paginas publicadas afirmam que «cada corte esta registrado em relacao anexa ao
produto». Essa relacao era `anexos/cortes.csv`, com 7 registros de 12/09,
enquanto as declaracoes chegaram a 27 em tres superficies. A afirmacao publicada
estava descoberta, e quem recebesse o produto nao tinha como auditar o que saiu.

## Por que ela fica FORA do git

Esta planilha guarda o TEXTO SUPRIMIDO. E o unico artefato do projeto que o
guarda, e por isso nao entra no repositorio, nao vai ao Hub e nao e publicada:
seria o mesmo que nao ter cortado. `codebook/supressoes.csv`, que e versionado,
guarda so o contexto, o comprimento e o `sha256` do trecho, nunca o trecho.

Ela segue para IPPLAM e CEPAL pelo canal do acervo, junto com as transcricoes de
entrega.

## De onde o texto e recuperado

Vinte das 27 declaracoes ainda nao estao aplicadas no acervo, e o trecho e lido
de la: do `Produto_4.docx`, do `dashboard.json` e das transcricoes do Anexo 05.
As outras sete ja foram aplicadas no Anexo 05 em 12/09, e por isso o texto delas
nao existe mais no acervo: vem de `anexos/cortes.csv`, que e o registro daquela
rodada e tambem fica fora do git.

## A trava

Cada trecho recuperado e conferido contra o `sha256` da declaracao. Divergencia
significa que o material mudou e que a relacao descreveria um corte que nao e o
que foi feito, entao a geracao para. Declaracao que nao possa ser recuperada de
nenhuma das duas fontes tambem para a geracao: uma relacao incompleta que se
apresenta como completa e pior do que nao ter relacao.
"""
from __future__ import annotations

import csv
import io
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from tools import supressoes as S                                # noqa: E402
from tools.acervo import fonte_transcricoes, SUFIXO              # noqa: E402

DESTINO = ROOT / "saida_entrega" / "Relacao_de_Cortes.xlsx"
CORTES_ANTIGOS = ROOT / "anexos" / "cortes.csv"
DASHBOARD = ROOT / "codebook" / "dashboard.json"
PRODUTO4 = ROOT / "anexos" / "Produto_4.docx"

TS = re.compile(r"\b(\d{2}:\d{2}:\d{2})\b")

ROTULO_FONTE = {"transcricao": "Transcrição", "produto4": "Produto 04",
                "evidencia": "Evidência no painel"}


def _efeito(s) -> str:
    if s.substituto == S.MARCA:
        return "corte marcado com (…)"
    if not s.substituto:
        return "remoção sem marca"
    return "substituição"


def _marca_tempo(paragrafos, idx: int) -> str:
    """A marca de tempo mais próxima antes do parágrafo."""
    for i in range(idx, -1, -1):
        m = TS.search(paragrafos[i])
        if m:
            return m.group(1)
    return ""


def _antigos() -> dict:
    if not CORTES_ANTIGOS.exists():
        return {}
    with io.open(CORTES_ANTIGOS, encoding="utf-8-sig", newline="") as f:
        return {(r["sessao"].strip(), int(r["paragrafo"])): r["suprimido"]
                for r in csv.DictReader(f)}


def coletar() -> list[dict]:
    """As 27 declarações com o trecho recuperado e conferido pelo hash."""
    import docx
    sups = S.carregar()
    antigos = _antigos()
    linhas, sem_texto, hash_errado = [], [], []

    def registrar(s, trecho, origem, local, ts="", extra=""):
        if trecho is None:
            sem_texto.append(f"{s.fonte}/{s.alvo} §{s.paragrafo}")
            return
        if S.sha(trecho) != s.sha256:
            hash_errado.append(f"{s.fonte}/{s.alvo} §{s.paragrafo}")
            return
        linhas.append({
            "fonte": ROTULO_FONTE.get(s.fonte, s.fonte), "alvo": s.alvo,
            "local": local, "marca_tempo": ts, "trecho": trecho,
            "substituto": s.substituto, "efeito": _efeito(s), "motivo": s.motivo,
            "n_chars": s.n_chars, "sha256": s.sha256, "origem_do_texto": origem,
            "observacao": extra})

    pars_p4 = [p.text for p in docx.Document(str(PRODUTO4)).paragraphs]
    for s in [x for x in sups if x.fonte == "produto4"]:
        a = S._casar(pars_p4[s.paragrafo], s) if s.paragrafo < len(pars_p4) else None
        registrar(s, pars_p4[s.paragrafo][a[0]:a[1]] if a else None,
                  "anexos/Produto_4.docx", f"parágrafo {s.paragrafo}")

    d = json.loads(io.open(DASHBOARD, encoding="utf-8").read())
    por = {}
    for e in d["evidence"]:
        por.setdefault(e["interview"], []).append(e)
    for s in [x for x in sups if x.fonte == "evidencia"]:
        fila = por.get(s.alvo, [])
        txt = fila[s.paragrafo]["excerpt"] if s.paragrafo < len(fila) else None
        dim = fila[s.paragrafo].get("dim", "") if s.paragrafo < len(fila) else ""
        if txt is not None:
            # aplica as declaracoes anteriores do mesmo trecho, na ordem do CSV
            for ant in sups:
                if ant is s:
                    break
                if (ant.fonte == "evidencia" and ant.alvo == s.alvo
                        and ant.paragrafo == s.paragrafo):
                    txt = S.aplicar_em(txt, ant)
        a = S._casar(txt, s) if txt is not None else None
        registrar(s, txt[a[0]:a[1]] if a else None, "codebook/dashboard.json",
                  f"evidência #{s.paragrafo}, dimensão {dim}")

    with fonte_transcricoes() as origem:
        if origem is None:
            raise SystemExit("acervo indisponível: a relação não pode ser conferida.")
        cache = {}
        for s in [x for x in sups if x.fonte == "transcricao"]:
            if s.alvo not in cache:
                caminho = Path(origem) / (s.alvo + SUFIXO)
                cache[s.alvo] = [p.text for p in docx.Document(str(caminho)).paragraphs]
            pars = cache[s.alvo]
            a = S._casar(pars[s.paragrafo], s) if s.paragrafo < len(pars) else None
            if a:
                registrar(s, pars[s.paragrafo][a[0]:a[1]], "Anexo 05",
                          f"parágrafo {s.paragrafo}", _marca_tempo(pars, s.paragrafo))
            else:
                registrar(s, antigos.get((s.alvo, s.paragrafo)), "anexos/cortes.csv",
                          f"parágrafo {s.paragrafo}", _marca_tempo(pars, s.paragrafo),
                          "já aplicado no Anexo 05 em 12/09")

    if sem_texto or hash_errado:
        raise SystemExit(
            "a relação sairia incompleta ou errada, e não foi gerada.\n"
            + ("  sem texto recuperável: " + ", ".join(sem_texto) + "\n" if sem_texto else "")
            + ("  hash divergente (o material mudou): " + ", ".join(hash_errado)
               if hash_errado else ""))
    if len(linhas) != len(sups):
        raise SystemExit(f"{len(linhas)} linhas para {len(sups)} declarações")
    return linhas


def gerar(destino: Path = DESTINO) -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    sys.path.insert(0, str(ROOT / "tools" / "hub"))
    from tools.hub import retencao as R

    linhas = coletar()
    mapa = R.carregar()

    wb = Workbook()
    cab_fill = PatternFill("solid", fgColor="0F4C5C")
    cab_font = Font(color="FFFFFF", bold=True, size=10)

    def folha(ws, cabecalho, dados, larguras):
        ws.append(cabecalho)
        for c in ws[1]:
            c.fill, c.font = cab_fill, cab_font
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for linha in dados:
            ws.append(linha)
        for j, w in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
        for linha in ws.iter_rows(min_row=2):
            for c in linha:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    ws = wb.active
    ws.title = "CORTES"
    folha(ws,
          ["Superfície", "Sessão ou documento", "Localização", "Marca de tempo",
           "Trecho suprimido", "O que entrou no lugar", "Efeito", "Motivo",
           "Caracteres", "sha256 do trecho", "Origem do texto", "Observação"],
          [[x["fonte"], x["alvo"], x["local"], x["marca_tempo"], x["trecho"],
            x["substituto"], x["efeito"], x["motivo"], x["n_chars"], x["sha256"],
            x["origem_do_texto"], x["observacao"]] for x in linhas],
          [16, 18, 22, 13, 62, 34, 20, 48, 11, 22, 22, 26])

    ws2 = folha(wb.create_sheet("SESSOES RETIDAS"),
                ["Sessão", "Publicada", "Termo próprio", "Passagens medidas",
                 "Limiar", "Critério", "Decidido em", "Decidido por", "Motivo"],
                [[c, "sim" if r["publica"] else "NÃO", r.get("termo", ""),
                  r["n_passagens"], r["limiar"], r.get("criterio", ""),
                  r.get("decidido_em", ""), r.get("decidido_por", ""),
                  r.get("motivo", "")]
                 for c, r in sorted(mapa.items())],
                [18, 11, 22, 17, 9, 46, 13, 22, 60])

    ws3 = wb.create_sheet("LEIA-ME", 0)
    for t in [
        "Relação de cortes por identificabilidade",
        "",
        "Consultoria técnica IPPLAM, CEPAL/ONU. Condições habilitantes ao "
        "financiamento climático urbano de Maringá. Anexo ao Produto 04.",
        "",
        "ESTE ARQUIVO NÃO PODE SER PUBLICADO.",
        "",
        "Ele guarda o texto que foi suprimido das transcrições, do relatório e do "
        "painel. É o único artefato do projeto que guarda esse texto. Publicá-lo "
        "desfaz todos os cortes de uma vez. Fica fora do repositório e fora do Hub, "
        "e circula apenas pelo canal do acervo, com as transcrições de entrega.",
        "",
        "O codebook versionado (codebook/supressoes.csv) registra cada corte pelo "
        "contexto, pelo comprimento e pelo sha256 do trecho, nunca pelo trecho. É "
        "assim que o corte pode ser auditado sem que o material protegido passe a "
        "viver no repositório.",
        "",
        "ABA «CORTES»",
        "Uma linha por supressão declarada, nas três superfícies em que o material "
        "circula: a transcrição anonimizada, o texto do Produto 04 e o trecho de "
        "evidência publicado no painel. A coluna «Efeito» diz o que aconteceu: corte "
        "marcado com (…), remoção sem marca quando o que identificava era só o "
        "marcador de primeira pessoa, ou substituição quando havia conteúdo a "
        "preservar.",
        "",
        "ABA «SESSÕES RETIDAS»",
        "A outra forma de omissão, e a maior. Em oito sessões o participante nomeia o "
        "próprio empregador tantas vezes que a supressão pontual não sobreviveria: o "
        "número de cortes necessário destruiria o conteúdo. Essas sessões saíram da "
        "camada pública e integram o anexo de entrega. A coluna «Passagens medidas» é "
        "recalculada a cada publicação e confrontada com o valor declarado.",
        "",
        "Todo o conteúdo é derivado de codebook/supressoes.csv, "
        "codebook/retencao_publica.csv e do acervo. Reproduzível por "
        "python -m tools.relacao_cortes.",
    ]:
        ws3.append([t])
    ws3.column_dimensions["A"].width = 110
    for linha in ws3.iter_rows():
        for c in linha:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws3["A1"].font = Font(bold=True, size=14, color="0F4C5C")
    ws3["A5"].font = Font(bold=True, size=11, color="B00020")

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(destino))
    return {"cortes": len(linhas), "sessoes": len(mapa),
            "retidas": sum(1 for r in mapa.values() if not r["publica"]),
            "destino": destino, "bytes": destino.stat().st_size}


def main() -> int:
    r = gerar()
    print(f"gerado: {r['destino']}")
    print(f"  {r['cortes']} cortes declarados, todos conferidos pelo sha256")
    print(f"  {r['sessoes']} sessões, {r['retidas']} retidas da camada pública")
    print(f"  {r['bytes'] // 1024} KB")
    print("  FORA do git: guarda o texto suprimido. Não publique.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
